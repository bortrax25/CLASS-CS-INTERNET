"""Tests de src/report.py: genera el Excel con datos sintéticos y lo vuelve a leer."""
from datetime import date

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src.report import COLUMNAS, exportar_excel, ruta_reporte

ENCABEZADOS = [c[0] for c in COLUMNAS]


def col(nombre):
    return ENCABEZADOS.index(nombre) + 1


def fila_metricas(ticker, pmin, pmax, revisar="", salto=5.0):
    return {
        "ticker": ticker, "precio_min": pmin, "fecha_min": pd.Timestamp("2026-01-05"),
        "precio_max": pmax, "fecha_max": pd.Timestamp("2026-06-25"),
        "rango_abs": pmax - pmin, "rango_pct": (pmax / pmin - 1) * 100, "ratio": pmax / pmin,
        "direccion": "subió", "precio_actual": pmax * 0.9, "fecha_actual": pd.Timestamp("2026-10-06"),
        "volumen_promedio": 1e6, "volumen_usd_promedio": 5e7, "dias_con_dato": 191,
        "cobertura": 1.0, "salto_max_pct": salto, "fecha_salto": pd.Timestamp("2026-03-02"),
        "moneda": "USD", "revisar": revisar,
    }


@pytest.fixture
def libro(tmp_path):
    pasan = pd.DataFrame([
        fila_metricas("AAA", 10.0, 85.0),
        fila_metricas("CTVA", 11.92, 90.32, revisar="salto de -83.8% el 2026-10-01", salto=-83.8),
        fila_metricas("BBB", 50.0, 100.0),
    ])
    excluidas = pd.DataFrame({"ticker": ["HONA", "ZZZ"], "motivo": ["baja cobertura (41% < 80%)", "sin datos"]})
    universo = pd.DataFrame({"ticker": ["AAA", "BBB", "CTVA", "HONA"],
                             "nombre": ["Alfa", "Beta", "Corteva", "Honeywell Aerospace"],
                             "sector": ["Tech", "Energy", "Materials", "Industrials"]})
    ruta = exportar_excel(pasan, excluidas, tmp_path / "x.xlsx", top_n=2, universo=universo,
                          parametros={"Universo": "sp500", "Inicio": date(2026, 1, 1)})
    return load_workbook(ruta)


def test_hojas(libro):
    assert libro.sheetnames == ["Top 2", "Todos", "Excluidos", "Parámetros"]


def test_top_y_todos(libro):
    top, todos = libro["Top 2"], libro["Todos"]
    assert [c.value for c in top[1]] == ENCABEZADOS
    assert top.max_row == 3 and todos.max_row == 4
    assert [todos.cell(r, col("Ticker")).value for r in (2, 3, 4)] == ["AAA", "CTVA", "BBB"]
    assert [todos.cell(r, col("#")).value for r in (2, 3, 4)] == [1, 2, 3]
    assert todos.cell(2, col("Nombre")).value == "Alfa"
    assert todos.cell(3, col("Sector")).value == "Materials"


def test_formulas_y_formatos(libro):
    ws = libro["Todos"]
    mn = ws.cell(2, col("Precio mín")).column_letter
    mx = ws.cell(2, col("Precio máx")).column_letter
    assert ws.cell(2, col("Rango %")).value == f"=IF(AND(COUNT({mn}2,{mx}2)=2,{mn}2>0),{mx}2/{mn}2-1,\"\")"
    assert ws.cell(2, col("Rango abs")).value.startswith("=IF(")
    assert ws.cell(2, col("Rango %")).number_format == "0.0%"
    assert ws.cell(2, col("Fecha mín")).number_format == "yyyy-mm-dd"
    assert ws.cell(2, col("Fecha mín")).value.date() == date(2026, 1, 5)
    # Los porcentajes se guardan como fracciones
    assert ws.cell(3, col("Salto máx 1 día")).value == pytest.approx(-0.838)
    assert ws.cell(2, col("Cobertura")).value == 1.0
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == f"A1:{ws.cell(1, len(ENCABEZADOS)).column_letter}4"
    assert ws.cell(1, 1).font.name == "Arial"


def test_resalta_las_que_hay_que_revisar(libro):
    ws = libro["Todos"]
    assert ws.cell(3, col("Revisar")).value.startswith("salto de -83.8%")
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith("FFF2CC")
    assert ws.cell(2, 1).fill.fgColor.rgb in (None, "00000000")
    assert ws.cell(2, col("Revisar")).value is None


def test_excluidos(libro):
    ws = libro["Excluidos"]
    filas = [[c.value for c in r] for r in ws.iter_rows(min_row=1)]
    assert filas == [["Ticker", "Nombre", "Motivo"],
                     ["HONA", "Honeywell Aerospace", "baja cobertura (41% < 80%)"],
                     ["ZZZ", None, "sin datos"]]
    assert ws.freeze_panes == "A2"


def test_parametros(libro):
    ws = libro["Parámetros"]
    assert ws["A1"].value == "Universo" and ws["B1"].value == "sp500"


def test_nan_queda_vacio(tmp_path):
    f = fila_metricas("A", 10.0, 20.0)
    f["salto_max_pct"], f["fecha_salto"] = np.nan, pd.NaT
    ruta = exportar_excel(pd.DataFrame([f]), pd.DataFrame(columns=["ticker", "motivo"]), tmp_path / "y.xlsx")
    ws = load_workbook(ruta)["Todos"]
    assert ws.cell(2, col("Salto máx 1 día")).value is None
    assert ws.cell(2, col("Fecha salto")).value is None
    assert load_workbook(ruta)["Excluidos"].max_row == 1


def test_ruta_reporte(tmp_path):
    assert ruta_reporte(tmp_path, "sp500", date(2026, 10, 6)).name == "top_sp500_2026-10-06.xlsx"
