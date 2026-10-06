"""Tests de main.py con la descarga sustituida por datos sintéticos (sin internet)."""
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

import main
from src.config import cargar_config


def precios_falsos(llamadas):
    def get_prices(tickers, start, end, **kw):
        llamadas.append({"start": start, "end": end, **kw})
        filas = []
        fechas = pd.bdate_range("2026-01-02", "2026-02-27")
        series = {"AAPL": (100, 1.0), "MSFT": (50, 3.0), "NVDA": (20, -0.2), "BRK-B": (0.5, 0.0)}
        for t, (base, paso) in series.items():
            for i, f in enumerate(fechas):
                c = base + paso * i
                filas.append({"ticker": t, "fecha": f, "open": c, "high": c * 1.02, "low": c * 0.98,
                              "close": c, "volume": 1e6})
        return pd.DataFrame(filas)
    return get_prices


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    cfg = cargar_config()
    cfg["periodo"] = {"inicio": "2026-01-01", "fin": "2026-02-27"}
    cfg["rutas"] = {"cache": str(tmp_path / "cache"), "output": str(tmp_path / "output")}
    llamadas = []
    monkeypatch.setattr(main, "get_prices", precios_falsos(llamadas))
    return cfg, llamadas, tmp_path


def test_corrida_completa(entorno, capsys):
    cfg, llamadas, tmp = entorno
    ruta = main.ejecutar("sp500", top=2, cfg=cfg)

    assert ruta == tmp / "output" / "top_sp500_2026-02-27.xlsx"
    assert llamadas[0]["start"] == "2026-01-01" and llamadas[0]["end"] == "2026-02-28"  # fin inclusivo
    assert llamadas[0]["ruta_fallidos"] == tmp / "output" / "fallidos_sp500.csv"
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Top 2", "Todos", "Excluidos", "Parámetros"]
    top = [wb["Top 2"].cell(r, 2).value for r in (2, 3)]
    assert top == ["MSFT", "NVDA"]  # 50→170 (+240 %) y 20→12 (66,7 %); AAPL 100→140 (40 %)
    excluidos = {r[0]: r[2] for r in wb["Excluidos"].iter_rows(min_row=2, values_only=True)}
    assert excluidos["BRK-B"].startswith("precio bajo")
    assert excluidos["AMZN"] == "sin datos"  # del universo, sin precios
    assert "MSFT" in capsys.readouterr().out


def test_sin_filtros_intraday_y_nombre(entorno):
    cfg, _, tmp = entorno
    ruta = main.ejecutar("sp500", sin_filtros=True, intraday=True, cfg=cfg)
    assert ruta.name == "top_sp500_2026-02-27_intraday_sinfiltros.xlsx"
    wb = load_workbook(ruta)
    tickers = [r[1] for r in wb["Todos"].iter_rows(min_row=2, values_only=True)]
    assert "BRK-B" in tickers  # sin filtros entra la penny
    params = {r[0]: r[1] for r in wb["Parámetros"].iter_rows(values_only=True)}
    assert params["Filtros"].startswith("ninguno")
    assert params["Precio"].startswith("intraday")


def test_hasta_cambia_solo_esa_corrida(entorno):
    cfg, llamadas, _ = entorno
    ruta = main.ejecutar("nasdaq100", hasta="2026-02-13", cfg=cfg)
    assert llamadas[0]["end"] == "2026-02-14"
    assert ruta.name == "top_nasdaq100_2026-02-13.xlsx"
    assert cfg["periodo"]["fin"] == "2026-02-27"  # no modifica la config recibida


def test_hasta_anterior_al_inicio(entorno):
    cfg, _, _ = entorno
    with pytest.raises(ValueError, match="anterior al inicio"):
        main.ejecutar("sp500", hasta="2025-12-31", cfg=cfg)


def test_sin_precios(entorno, monkeypatch):
    cfg, _, _ = entorno
    monkeypatch.setattr(main, "get_prices", lambda *a, **k: pd.DataFrame())
    with pytest.raises(RuntimeError):
        main.ejecutar("sp500", cfg=cfg)


def test_nombre_archivo():
    assert main.nombre_archivo("sp500", date(2026, 10, 5), False, False) == "top_sp500_2026-10-05.xlsx"


@pytest.mark.parametrize("args", [
    ["--universe", "bvl"],
    ["--universe", "sp500", "--top", "0"],
    ["--universe", "sp500", "--hasta", "5-10-2026"],
    [],
])
def test_argumentos_invalidos(args):
    with pytest.raises(SystemExit) as e:
        main.main(args)
    assert e.value.code == 2
