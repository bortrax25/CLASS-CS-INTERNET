"""Exportación del ranking a Excel.

Archivo output/top_{universo}_{fecha}.xlsx con las hojas:
1. "Top N": las N primeras por rango %
2. "Todos": todas las que pasaron los filtros
3. "Excluidos": ticker, nombre y motivo
4. "Parámetros": periodo, tipo de precio, filtros y fuente (para saber cómo se generó)

Rango abs, rango % y ratio son fórmulas sobre el precio mínimo y máximo. Los porcentajes
se guardan como fracciones con formato de porcentaje. Las filas con un salto de un día
sospechoso (columna "Revisar") van resaltadas.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FUENTE = "Arial"
F_NORMAL = Font(name=FUENTE, size=10)
F_TITULO = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
RELLENO_TITULO = PatternFill("solid", fgColor="1F4E78")
RELLENO_REVISAR = PatternFill("solid", fgColor="FFF2CC")

FMT_PRECIO = "#,##0.00"
FMT_ENTERO = "#,##0"
FMT_PCT = "0.0%"
FMT_FECHA = "yyyy-mm-dd"

# (encabezado, clave en el DataFrame, formato, ancho, comentario)
# Las claves rango_abs, rango_pct y ratio se escriben como fórmulas.
COLUMNAS = [
    ("#", "_posicion", "0", 5, None),
    ("Ticker", "ticker", None, 9, None),
    ("Nombre", "nombre", None, 30, None),
    ("Sector", "sector", None, 22, "S&P 500: sector GICS. Nasdaq 100: industria ICB."),
    ("Precio mín", "precio_min", FMT_PRECIO, 11, None),
    ("Fecha mín", "fecha_min", FMT_FECHA, 11, None),
    ("Precio máx", "precio_max", FMT_PRECIO, 11, None),
    ("Fecha máx", "fecha_max", FMT_FECHA, 11, None),
    ("Rango abs", "rango_abs", FMT_PRECIO, 11, "Precio máx - precio mín, en la moneda de la acción."),
    ("Rango %", "rango_pct", FMT_PCT, 10, "(Precio máx - precio mín) / precio mín. Ordena el ranking."),
    ("Ratio", "ratio", '0.00"x"', 8, "Precio máx / precio mín."),
    ("Dirección", "direccion", None, 11,
     "subió: el mínimo fue antes que el máximo. cayó: al revés. sin cambio: el mismo día."),
    ("Precio actual", "precio_actual", FMT_PRECIO, 12, None),
    ("Fecha actual", "fecha_actual", FMT_FECHA, 11, None),
    ("Volumen prom (acciones)", "volumen_promedio", FMT_ENTERO, 14, None),
    ("Volumen prom (USD)", "volumen_usd_promedio", FMT_ENTERO, 16,
     "Promedio diario de cierre x volumen. Es la medida de liquidez del filtro."),
    ("Cobertura", "cobertura", "0%", 10, "Días con dato / sesiones de mercado del periodo."),
    ("Salto máx 1 día", "salto_max_pct", "+0.0%;-0.0%;0.0%", 12, "Mayor variación de cierre a cierre, con signo."),
    ("Fecha salto", "fecha_salto", FMT_FECHA, 11, None),
    ("Moneda", "moneda", None, 8, None),
    ("Revisar", "revisar", None, 55,
     "Salto de un día mayor al umbral: puede ser un spin-off o split sin ajustar. No se excluye."),
]
DIVIDIR_100 = {"salto_max_pct"}  # vienen en puntos porcentuales; Excel quiere fracciones


def _encabezados(ws: Worksheet, titulos: list[tuple[str, int, str | None]]) -> None:
    for c, (titulo, ancho, comentario) in enumerate(titulos, 1):
        celda = ws.cell(row=1, column=c, value=titulo)
        celda.font, celda.fill = F_TITULO, RELLENO_TITULO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if comentario:
            celda.comment = Comment(comentario, "rango-acciones")
        ws.column_dimensions[get_column_letter(c)].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"  # encabezados y ticker siempre visibles


def _valor(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    return v.item() if hasattr(v, "item") else v


def _hoja_ranking(ws: Worksheet, df: pd.DataFrame) -> None:
    _encabezados(ws, [(t, a, com) for t, _, _, a, com in COLUMNAS])
    letra = {clave: get_column_letter(i) for i, (_, clave, _, _, _) in enumerate(COLUMNAS, 1)}
    for fila_n, (pos, fila) in enumerate(df.iterrows(), start=2):
        mn, mx = f"{letra['precio_min']}{fila_n}", f"{letra['precio_max']}{fila_n}"
        formulas = {
            "rango_abs": f"=IF(COUNT({mn},{mx})=2,{mx}-{mn},\"\")",
            "rango_pct": f"=IF(AND(COUNT({mn},{mx})=2,{mn}>0),{mx}/{mn}-1,\"\")",
            "ratio": f"=IF(AND(COUNT({mn},{mx})=2,{mn}>0),{mx}/{mn},\"\")",
        }
        revisar = bool(fila.get("revisar"))
        for col, (_, clave, fmt, _, _) in enumerate(COLUMNAS, 1):
            if clave == "_posicion":
                v = fila_n - 1
            elif clave in formulas:
                v = formulas[clave]
            else:
                v = _valor(fila.get(clave))
                if clave in DIVIDIR_100 and v is not None:
                    v = v / 100
            celda = ws.cell(row=fila_n, column=col, value=v)
            celda.font = F_NORMAL
            if fmt:
                celda.number_format = fmt
            if revisar:
                celda.fill = RELLENO_REVISAR
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(df) + 1}"


def _hoja_excluidos(ws: Worksheet, excluidas: pd.DataFrame) -> None:
    _encabezados(ws, [("Ticker", 9, None), ("Nombre", 30, None), ("Motivo", 70, None)])
    ws.freeze_panes = "A2"
    for i, fila in enumerate(excluidas.itertuples(index=False), start=2):
        for c, v in enumerate([fila.ticker, getattr(fila, "nombre", None), fila.motivo], 1):
            ws.cell(row=i, column=c, value=_valor(v)).font = F_NORMAL
    if len(excluidas):
        ws.auto_filter.ref = f"A1:C{len(excluidas) + 1}"


def _hoja_parametros(ws: Worksheet, parametros: dict) -> None:
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    for i, (k, v) in enumerate(parametros.items(), start=1):
        ws.cell(row=i, column=1, value=k).font = F_NEGRITA
        ws.cell(row=i, column=2, value=_valor(v)).font = F_NORMAL


def exportar_excel(
    pasan: pd.DataFrame,
    excluidas: pd.DataFrame,
    ruta: Path | str,
    *,
    top_n: int = 50,
    universo: pd.DataFrame | None = None,
    parametros: dict | None = None,
) -> Path:
    """Escribe el Excel. `pasan` debe venir ordenado por rango %. `universo` (ticker, nombre,
    sector) añade el nombre y el sector de cada empresa."""
    if universo is not None:
        info = universo[["ticker", "nombre", "sector"]]
        pasan = pasan.drop(columns=["nombre", "sector"], errors="ignore").merge(info, on="ticker", how="left")
        excluidas = excluidas.drop(columns=["nombre"], errors="ignore").merge(
            info[["ticker", "nombre"]], on="ticker", how="left")

    wb = Workbook()
    ws_top = wb.active
    ws_top.title = f"Top {top_n}"
    _hoja_ranking(ws_top, pasan.head(top_n))
    _hoja_ranking(wb.create_sheet("Todos"), pasan)
    _hoja_excluidos(wb.create_sheet("Excluidos"), excluidas)
    if parametros:
        _hoja_parametros(wb.create_sheet("Parámetros"), parametros)
    wb.calculation.fullCalcOnLoad = True  # que Excel calcule las fórmulas al abrir

    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def ruta_reporte(dir_output: Path | str, universo: str, fecha: date | None = None) -> Path:
    return Path(dir_output) / f"top_{universo}_{(fecha or date.today()).isoformat()}.xlsx"
